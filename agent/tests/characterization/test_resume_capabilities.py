from io import BytesIO

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from agent_service.application.use_cases.parse_resume import ParseResumeUseCase
from agent_service.application.use_cases.analyze_jd import AnalyzeJDUseCase
from agent_service.domain.models import JobDescription, PersonalInfo, Resume
from agent_service.infra.storage.files import FileStorage


class CapturingExtractor:
    def __init__(self):
        self.text = None

    async def extract_resume_data(self, text):
        self.text = text
        return Resume(personal_info=PersonalInfo(fullname="Ada"), raw_text=text)


def _docx_bytes():
    stream = BytesIO()
    document = Document()
    document.add_paragraph("  Ada   Lovelace  ")
    document.add_paragraph("Python\tEngineer")
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes():
    stream = BytesIO()
    page = canvas.Canvas(stream)
    page.drawString(72, 720, "Ada   Lovelace")
    page.drawString(72, 700, "Python Engineer")
    page.save()
    return stream.getvalue()


@pytest.mark.parametrize(
    ("filename", "content"),
    [("resume.docx", _docx_bytes()), ("resume.pdf", _pdf_bytes())],
    ids=["docx", "pdf"],
)
async def test_parse_resume_normalizes_extracted_text(filename, content, tmp_path):
    extractor = CapturingExtractor()
    use_case = ParseResumeUseCase(extractor, FileStorage(str(tmp_path)))

    result = await use_case.execute(content, filename)

    assert extractor.text == "Ada Lovelace\nPython Engineer"
    assert result.raw_text == extractor.text


@pytest.mark.parametrize(
    ("filename", "content"),
    [("resume.pdf", b"not a pdf"), ("resume.docx", b"not a docx")],
)
async def test_parse_resume_rejects_malformed_files(filename, content, tmp_path):
    use_case = ParseResumeUseCase(CapturingExtractor(), FileStorage(str(tmp_path)))

    with pytest.raises(ValueError, match="Failed to parse"):
        await use_case.execute(content, filename)


class CapturingJDAnalyzer:
    def __init__(self):
        self.text = None

    async def analyze(self, text):
        self.text = text
        return JobDescription(raw_text=text, position="Engineer", required_skills=["Python"])


async def test_jd_analysis_preserves_a_stable_normalized_source():
    analyzer = CapturingJDAnalyzer()
    use_case = AnalyzeJDUseCase(analyzer)

    result = await use_case.execute("  " + ("Python platform engineer. " * 3) + "  ")

    assert analyzer.text == "Python platform engineer. Python platform engineer. Python platform engineer."
    assert result.raw_text == analyzer.text


async def test_jd_analysis_is_stable_across_equivalent_whitespace():
    analyzer = CapturingJDAnalyzer()
    use_case = AnalyzeJDUseCase(analyzer)

    first = await use_case.execute("Python  platform\nengineer. " * 3)
    second = await use_case.execute("  Python platform engineer.\t" * 3)

    assert first.raw_text == second.raw_text
    assert first.position == second.position
    assert first.required_skills == second.required_skills
