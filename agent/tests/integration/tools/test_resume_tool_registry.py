from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from agent_service.application.use_cases.export_resume import ExportResumeUseCase
from agent_service.domain.models import (
    BulletOptimization,
    PersonalInfo,
    TailoredResume,
    Resume,
    JobDescription,
)
from agent_service.infra.generators.pdf_generator import PDFGenerator
from agent_service.infra.generators.word_generator import WordGenerator
from agent_service.tool_registry import ToolNotAllowedError
from agent_service.tools.resume_tools import build_resume_tool_registry


class Unused:
    async def execute(self, *args, **kwargs):
        raise AssertionError("unused capability called")


class TailorStub:
    async def execute(self, master, jd):
        return TailoredResume(
            master_resume_id=master.id,
            jd_id=jd.id,
            personal_info=PersonalInfo(fullname="Ada Lovelace"),
            selected_experience_ids=["exp-1"],
            selected_bullet_optimizations=[
                BulletOptimization(
                    bullet_id="bullet-1",
                    original_text="Built an analytical engine",
                    optimized_text="Engineered an analytical engine with Python",
                    source_evidence=["Built an analytical engine"],
                )
            ],
            selected_education_ids=["edu-1"],
            selected_skills=["Python"],
        )


class RetrievalStub:
    def __init__(self):
        self.indexed = []

    async def index(self, **kwargs):
        self.indexed.append(kwargs)

    async def retrieve(self, query, *, owner, run_id, top_k):
        assert owner == "user-a"
        assert run_id == "run-1"
        return [
            {
                "content": "Never invent unsupported metrics.",
                "source_type": "knowledge",
                "source_id": "anti_hallucination_rules.json",
                "score": 0.91,
                "metadata": {"scope": "global"},
            }
        ]


class ParseStub:
    async def execute(self, content, filename):
        return Resume(raw_text="Owned Python resume evidence")


class AnalyzeStub:
    async def execute(self, text):
        return JobDescription(raw_text=text)


def _registry():
    return build_resume_tool_registry(
        parse_resume=Unused(),
        analyze_jd=Unused(),
        tailor_resume=Unused(),
        export_resume=ExportResumeUseCase(PDFGenerator(), WordGenerator()),
        evidence_retriever=RetrievalStub(),
    )


def _resume():
    return {
        "personal_info": {
            "fullname": "Ada Lovelace",
            "email": "ada@example.com",
            "phone": "+1 555 0100",
            "location": "London",
        },
        "summary": "Computing pioneer",
        "experiences": [],
        "education": [],
        "skills": [{"name": "Python", "category": "language"}],
    }


def test_resume_registry_publishes_exact_typed_contracts():
    metadata = {item["name"]: item for item in _registry().describe()}

    assert set(metadata) == {
        "parse_resume",
        "analyze_jd",
        "propose_tailoring",
        "export_approved_resume",
        "retrieve_evidence",
    }
    expected = {
        "parse_resume": ("read", ["awaiting_resume"], {"filename", "content"}, "ParseResumeResult"),
        "analyze_jd": ("read", ["awaiting_jd", "ready"], {"text"}, "AnalyzeJDResult"),
        "propose_tailoring": (
            "proposal",
            ["analyzing", "proposal_ready"],
            {"master_resume", "job_description"},
            "TailorResumeResult",
        ),
        "export_approved_resume": (
            "mutation",
            ["approved"],
            {"resume", "format"},
            "ExportResumeResult",
        ),
        "retrieve_evidence": (
            "read",
            ["analyzing", "proposal_ready"],
            {"query"},
            "RetrieveEvidenceResult",
        ),
    }
    for name, (effect, states, required, result_type) in expected.items():
        assert metadata[name]["effect"] == effect
        assert metadata[name]["allowed_states"] == states
        assert set(metadata[name]["input_schema"]["required"]) == required
        assert metadata[name]["result_type"] == result_type

    format_schema = metadata["export_approved_resume"]["input_schema"]["properties"]["format"]
    assert format_schema["enum"] == ["pdf", "docx"]


async def test_retrieval_uses_trusted_owner_context_and_returns_source_locators():
    result = await _registry().execute(
        "retrieve_evidence",
        {"query": "Can I add a metric?"},
        state="analyzing",
        owner="user-a",
        run_id="run-1",
    )

    assert result.items[0].source_id == "anti_hallucination_rules.json"
    assert result.items[0].score == 0.91


async def test_parsed_resume_is_indexed_under_trusted_owner():
    retrieval = RetrievalStub()
    registry = build_resume_tool_registry(
        parse_resume=ParseStub(), analyze_jd=Unused(), tailor_resume=Unused(),
        export_resume=ExportResumeUseCase(PDFGenerator(), WordGenerator()),
        evidence_retriever=retrieval,
    )

    await registry.execute(
        "parse_resume",
        {"filename": "resume.pdf", "content": b"resume"},
        state="awaiting_resume",
        owner="user-a",
        run_id="run-1",
    )

    assert retrieval.indexed == [{
        "owner": "user-a", "run_id": "run-1",
        "source_type": "resume",
        "source_id": "resume.pdf",
        "content": "Owned Python resume evidence",
    }]


async def test_analyzed_jd_is_indexed_under_trusted_owner():
    retrieval = RetrievalStub()
    registry = build_resume_tool_registry(
        parse_resume=Unused(), analyze_jd=AnalyzeStub(), tailor_resume=Unused(),
        export_resume=ExportResumeUseCase(PDFGenerator(), WordGenerator()),
        evidence_retriever=retrieval,
    )

    text = "Python platform engineer with distributed systems experience required."
    await registry.execute(
        "analyze_jd", {"text": text}, state="awaiting_jd", owner="user-a",
        run_id="run-1",
    )

    assert retrieval.indexed[0]["owner"] == "user-a"
    assert retrieval.indexed[0]["run_id"] == "run-1"
    assert retrieval.indexed[0]["source_type"] == "job_description"
    assert retrieval.indexed[0]["content"] == text




async def test_approved_resume_exports_valid_pdf_through_real_registry():
    result = await _registry().execute(
        "export_approved_resume",
        {"resume": _resume(), "format": "pdf"},
        state="approved",
    )

    assert result.content.startswith(b"%PDF")
    assert len(PdfReader(BytesIO(result.content)).pages) == 1


async def test_approved_resume_exports_valid_docx_through_real_registry():
    result = await _registry().execute(
        "export_approved_resume",
        {"resume": _resume(), "format": "docx"},
        state="approved",
    )

    document = Document(BytesIO(result.content))
    assert "Ada Lovelace" in "\n".join(p.text for p in document.paragraphs)


async def test_export_tool_cannot_run_before_approval():
    with pytest.raises(ToolNotAllowedError):
        await _registry().execute(
            "export_approved_resume",
            {"resume": _resume(), "format": "pdf"},
            state="reviewing_proposal",
        )


@pytest.mark.parametrize("format", ["pdf", "docx"])
async def test_tailoring_result_exports_selected_approved_content(format):
    registry = build_resume_tool_registry(
        parse_resume=Unused(),
        analyze_jd=Unused(),
        tailor_resume=TailorStub(),
        export_resume=ExportResumeUseCase(PDFGenerator(), WordGenerator()),
        evidence_retriever=RetrievalStub(),
    )
    proposal = await registry.execute(
        "propose_tailoring",
        {
            "master_resume": {
                "id": "master-1",
                "personal_info": {"fullname": "Ada Lovelace"},
                "experiences": [
                    {
                        "id": "exp-1",
                        "company": "Difference Engine Co",
                        "title": "Programmer",
                        "bullets": [
                            {"id": "bullet-1", "text": "Built an analytical engine"}
                        ],
                    },
                    {"id": "exp-2", "company": "Unselected Co", "bullets": []},
                ],
                "education": [
                    {"id": "edu-1", "school": "University of London", "degree": "Mathematics"}
                ],
                "skills": [
                    {"name": "Python", "category": "language"},
                    {"name": "Unselected Skill", "category": "other"},
                ],
            },
            "job_description": {"id": "jd-1", "raw_text": "Python engineer"},
        },
        state="analyzing",
    )
    exported = await registry.execute(
        "export_approved_resume",
        {"resume": proposal.tailored_resume, "format": format},
        state="approved",
    )

    if format == "pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(exported.content)).pages)
    else:
        document = Document(BytesIO(exported.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Difference Engine Co" in text
    assert "Programmer" in text
    assert "Engineered an analytical engine with Python" in text
    assert "University of London" in text
    assert "Python" in text
    assert "Unselected Co" not in text
